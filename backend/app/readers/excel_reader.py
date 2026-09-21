"""Excel reader (ARCHITECTURE.md §4.3/§4.5): one dataset per sheet."""

from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl

from ..config import settings
from .base import DiscoveredEntity
from .inference import infer_schema


def _cell_str(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value)


def discover(path: Path) -> list[DiscoveredEntity]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    entities: list[DiscoveredEntity] = []
    try:
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows_iter = ws.iter_rows(values_only=True)
            try:
                header = next(rows_iter)
            except StopIteration:
                continue
            columns = [
                (str(c).strip() if c is not None else f"col_{i + 1}")
                for i, c in enumerate(header)
            ]
            sample_rows: list[dict[str, str | None]] = []
            for i, row in enumerate(rows_iter):
                if i >= settings.schema_sample_rows:
                    break
                if all(v is None for v in row):
                    continue
                sample_rows.append(
                    {columns[j]: _cell_str(row[j]) for j in range(len(columns)) if j < len(row)}
                )
            schema = infer_schema(sample_rows, columns)
            entities.append(
                DiscoveredEntity(
                    entity_name=sheet_name,
                    format="excel",
                    spec={"sheet": sheet_name, "columns": columns},
                    columns=schema,
                    sample_row_count=len(sample_rows),
                )
            )
    finally:
        wb.close()
    return entities


def read_rows(path: Path, spec: dict[str, Any]):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[spec["sheet"]]
        columns: list[str] = spec["columns"]
        rows_iter = ws.iter_rows(values_only=True)
        next(rows_iter, None)  # header
        for row in rows_iter:
            if all(v is None for v in row):
                continue
            yield {columns[j]: _cell_str(row[j]) for j in range(len(columns)) if j < len(row)}
    finally:
        wb.close()
